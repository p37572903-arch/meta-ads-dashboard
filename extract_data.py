"""
Extracts the dashboard's row-level dataset from a Meta Ads Manager export.

Handles two cases automatically:

1. RAW export — a workbook with a sheet (typically named "Sheet1") holding
   Meta's full raw column set (~66 columns, one row per ad per date range:
   Campaign name, Ad set name, Ad name, Amount spent (INR), Purchases
   conversion value, Reach, Impressions, Link clicks, Website landing page
   views, Website adds to cart, Purchases, ThruPlays, Hook Rate, ATC to CI %,
   etc.). Two fields the dashboard needs aren't in the raw export and are
   reconstructed here:

     3 sec (3-second views)   = Hook Rate * Impressions
     CI (Checkouts Initiated) = ATC to CI % * Website adds to cart
                                  (the WEBSITE adds-to-cart column
                                   specifically — there is a separate, more
                                   general "Adds to cart" column that gives
                                   slightly different numbers; only "Website
                                   adds to cart" reproduces the correct CI).

   These formulas were reverse-engineered and verified row-for-row (535/535,
   bit-for-bit) against a hand-built reference sheet that had 3 sec/CI
   pre-computed, so they can be trusted as the correct reconstruction.

2. PRE-COMPUTED export — a workbook that already has a sheet with "3 sec" and
   "CI" columns (e.g. a hand-built "Sheet 3" some users maintain). That sheet
   is used as-is, no reconstruction needed — this keeps the script working
   whether the user hands over the raw export or their own simplified sheet.

Usage:
    python3 extract_data.py <path-to-xlsx> [output.json] [--sheet SHEET_NAME]

If --sheet is omitted, the script looks for a sheet with pre-computed 3 sec/CI
columns first; otherwise it uses "Sheet1" if present, or the first sheet whose
headers look like the raw export.
"""
import sys
import json
import argparse
import openpyxl

# Columns required to compute everything the dashboard needs from a RAW export.
RAW_REQUIRED = [
    "Campaign name", "Ad set name", "Ad name", "Amount spent (INR)",
    "Purchases conversion value", "Reach", "Impressions", "Link clicks",
    "Website landing page views", "Website adds to cart", "Purchases",
    "ThruPlays", "Hook Rate", "ATC to CI %",
]
# Presence of both means the sheet already has the derived columns pre-computed.
PRECOMPUTED_MARKERS = ["3 sec", "CI"]
# Optional per-row date, tried in this order — a day-broken-out Meta export has
# "Day"; a date-range export has "Reporting starts"/"Reporting ends" (same
# value per row for a single-day breakdown). Not required: a sheet with none
# of these still extracts fine, just without the dashboard's date filter.
DATE_COLUMN_CANDIDATES = ["Day", "Date", "Reporting starts", "Reporting ends"]


def _headers_of(ws):
    return {ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)}


def find_sheet(wb, requested=None):
    if requested:
        if requested not in wb.sheetnames:
            raise ValueError(f"Sheet '{requested}' not found. Available sheets: {wb.sheetnames}")
        return wb[requested]

    # Prefer a sheet that already has the derived columns (a hand-built "Sheet 3").
    for name in wb.sheetnames:
        headers = _headers_of(wb[name])
        if all(m in headers for m in PRECOMPUTED_MARKERS) and "Campaign name" in headers:
            return wb[name]

    # Otherwise this is a raw export: literally "Sheet1" if present, else the
    # first sheet whose headers match the raw column set.
    if "Sheet1" in wb.sheetnames:
        return wb["Sheet1"]
    for name in wb.sheetnames:
        headers = _headers_of(wb[name])
        if all(m in headers for m in RAW_REQUIRED):
            return wb[name]

    raise ValueError(
        "Could not find a usable sheet (no pre-computed 3 sec/CI columns, no "
        f"'Sheet1', and no sheet with all raw columns present). Available sheets: "
        f"{wb.sheetnames}. Required raw columns: {RAW_REQUIRED}"
    )


def extract(path, sheet_name=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = find_sheet(wb, sheet_name)
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    is_precomputed = all(m in headers for m in PRECOMPUTED_MARKERS)

    missing = [c for c in RAW_REQUIRED if c not in headers] if not is_precomputed else []
    if missing:
        raise ValueError(f"Sheet '{ws.title}' is missing required raw columns: {missing}")

    def cell(row, name):
        col = headers.get(name)
        return ws.cell(row=row, column=col).value if col else None

    def num(v):
        return float(v) if v is not None else 0.0

    date_col = next((c for c in DATE_COLUMN_CANDIDATES if c in headers), None)

    def date_str(v):
        if v is None:
            return None
        if hasattr(v, "strftime"):
            return v.strftime("%Y-%m-%d")
        s = str(v).strip()
        return s[:10] if s else None

    rows = []
    for r in range(2, ws.max_row + 1):
        ad_name = cell(r, "Ad name")
        if ad_name is None:
            continue

        if is_precomputed:
            sec3 = num(cell(r, "3 sec"))
            ci = num(cell(r, "CI"))
        else:
            hook = cell(r, "Hook Rate")
            impr = num(cell(r, "Impressions"))
            atc_ci_pct = cell(r, "ATC to CI %")
            web_atc = num(cell(r, "Website adds to cart"))
            sec3 = (num(hook) * impr) if hook is not None else 0.0
            ci = (num(atc_ci_pct) * web_atc) if atc_ci_pct is not None else 0.0

        rows.append({
            "c": cell(r, "Campaign name"),
            "s": cell(r, "Ad set name"),
            "a": ad_name,
            "spend": num(cell(r, "Amount spent (INR)")),
            "revenue": num(cell(r, "Purchases conversion value")),
            "reach": num(cell(r, "Reach")),
            "impr": num(cell(r, "Impressions")),
            "clicks": num(cell(r, "Link clicks")),
            "lpv": num(cell(r, "Website landing page views")),
            "atc": web_atc if not is_precomputed else num(cell(r, "Website adds to cart")),
            "purchases": num(cell(r, "Purchases")),
            "sec3": sec3,
            "thru": num(cell(r, "ThruPlays")),
            "ci": ci,
            "date": date_str(cell(r, date_col)) if date_col else None,
        })
    return rows, is_precomputed, ws.title, date_col


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", help="Path to the Meta Ads export .xlsx")
    parser.add_argument("output", nargs="?", default="data.json", help="Output JSON path (default: data.json)")
    parser.add_argument("--sheet", default=None, help="Force a specific sheet name instead of auto-detecting")
    args = parser.parse_args()

    rows, was_precomputed, sheet_title, date_col = extract(args.input, args.sheet)
    with open(args.output, "w") as f:
        json.dump(rows, f, separators=(",", ":"))

    mode = "pre-computed sheet (used as-is)" if was_precomputed else "raw export (3 sec / CI reconstructed)"
    date_note = f"date column '{date_col}'" if date_col else "no date column found — date filter will be disabled"
    print(f"Extracted {len(rows)} rows from sheet '{sheet_title}' [{mode}] [{date_note}] -> {args.output}")


if __name__ == "__main__":
    main()
